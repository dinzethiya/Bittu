#!/usr/bin/env python3
"""
Regenerate data/products.json from the AVC "AI Product Master" spreadsheet.

Usage:
    pip install openpyxl
    python3 tools/export_products.py "AVC AI Product Master - Enriched v2.xlsx"

Run this any time the spreadsheet changes (new products, prices, stock),
then commit the updated data/products.json — the live site picks it up
on the next page load, no other changes needed.
"""
import sys
import json
import os

try:
    import openpyxl
except ImportError:
    sys.exit("Missing dependency. Run:  pip install openpyxl")

CATEGORY_MAP = {
    "ACCESSORIES": "Accessories",
    "Accessories": "Accessories",
    "Bras": "Bras",
    "Brief": "Briefs & Trunks",
    "Camisoles": "Camisoles & Slips",
    "Kidswear": "Kidswear",
    "Men's Outerwear": "Men's Outerwear",
    "Men's Thermalwear": "Thermal Wear",
    "Thermal wear": "Thermal Wear",
    "NIGHTSUIT": "Nightsuits",
    "NIGHTY": "Nighties",
    "Panties": "Panties",
    "SHAPEWEAR": "Shapewear",
    "Shapewear": "Shapewear",
    "Tights": "Tights",
    "Vest": "Vests",
    "Women's Outerwear": "Women's Outerwear",
}


def clean(v):
    return "" if v is None else str(v).strip()


def yesno(v):
    return clean(v).lower() == "yes"


def norm_category(c):
    c = clean(c)
    return CATEGORY_MAP.get(c, c.title() if c else "Other")


def to_int(v):
    try:
        return int(v) if v not in (None, "") else None
    except (ValueError, TypeError):
        return None


def main():
    if len(sys.argv) < 2:
        sys.exit("Usage: python3 export_products.py <path-to-xlsx>")

    xlsx_path = sys.argv[1]
    if not os.path.exists(xlsx_path):
        sys.exit(f"File not found: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "AI Product Master" not in wb.sheetnames:
        sys.exit('Sheet "AI Product Master" not found in this workbook.')
    ws = wb["AI Product Master"]

    products = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        (pid, brand, model, category, ptype, gender, size_range, cup_range,
         padded, sports, strapless, feeding, shapewear, intent, keywords,
         fabric, wired, coverage, padding_level, style, colours, mrp,
         selling_price, stock_status, stock_qty, img_url, description,
         features, ideal_for, *_rest) = row + (None,) * (34 - len(row))

        mrp_val = to_int(mrp)
        price_val = to_int(selling_price) or mrp_val

        products.append({
            "id": clean(pid),
            "brand": clean(brand),
            "model": clean(model),
            "category": norm_category(category),
            "type": clean(ptype),
            "gender": clean(gender),
            "sizeRange": clean(size_range),
            "cupRange": clean(cup_range),
            "padded": yesno(padded),
            "sports": yesno(sports),
            "strapless": yesno(strapless),
            "feeding": yesno(feeding),
            "shapewear": yesno(shapewear),
            "intent": clean(intent),
            "fabric": clean(fabric),
            "wired": clean(wired),
            "coverage": clean(coverage),
            "paddingLevel": clean(padding_level),
            "style": clean(style),
            "colours": [c.strip() for c in clean(colours).split(";") if c.strip()],
            "mrp": mrp_val,
            "price": price_val,
            "stock": clean(stock_status) or "Unknown",
            "description": clean(description),
            "features": [f.strip() for f in clean(features).split(";") if f.strip()],
            "idealFor": clean(ideal_for),
        })

    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "products.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(products, f, ensure_ascii=False, separators=(",", ":"))

    print(f"Wrote {len(products)} products to {out_path}")


if __name__ == "__main__":
    main()
